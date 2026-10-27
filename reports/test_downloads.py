import io
import zipfile

from core.models import AuditLog
from core.testing import BaseTestCase
from students.models import Student

URL = "/api/downloads/"


class DownloadScopeTests(BaseTestCase):
    def get(self, what, query=""):
        return self.client.get(f"{URL}{what}/?export=csv{query}")

    def text(self, r):
        return b"".join(r.streaming_content).decode("utf-8") if r.streaming else r.content.decode("utf-8")

    def test_admin_gets_every_student_hod_only_own_department(self):
        self.as_user(self.admin)
        r = self.get("students")
        self.assertEqual(r.status_code, 200)
        t = self.text(r)
        self.assertTrue(t.startswith("﻿"))
        self.assertIn("Math Student 1", t)
        self.assertIn("Student 5", t)
        self.as_user(self.hod_cs)
        t = self.text(self.get("students"))
        self.assertIn("Student 5", t)
        self.assertNotIn("Math Student", t)

    def test_teacher_gets_only_own_classes_students(self):
        self.as_user(self.t1)
        t = self.text(self.get("students"))
        self.assertIn("Student 1", t)
        self.assertNotIn("Math Student", t)
        self.as_user(self.t_math)
        t = self.text(self.get("students"))
        self.assertIn("Math Student 1", t)
        self.assertNotIn("Student 5", t)

    def test_teacher_cannot_download_staff_or_hods(self):
        self.as_user(self.t1)
        self.assertEqual(self.get("staff").status_code, 403)
        self.assertEqual(self.get("hods").status_code, 403)
        self.assertEqual(self.client.get(URL + "all/").status_code, 403)

    def test_hod_gets_own_staff_only_and_no_hods_or_all(self):
        self.as_user(self.hod_cs)
        t = self.text(self.get("staff"))
        self.assertIn("prof_a", t)
        self.assertNotIn("prof_math", t)
        self.assertNotIn("hod_math", t)
        self.assertEqual(self.get("hods").status_code, 403)
        self.assertEqual(self.client.get(URL + "all/").status_code, 403)

    def test_admin_gets_hods_and_staff_of_everyone(self):
        self.as_user(self.admin)
        h = self.text(self.get("hods"))
        self.assertIn("hod_cs", h)
        self.assertIn("hod_math", h)
        self.assertNotIn("prof_a", h)
        s = self.text(self.get("staff"))
        self.assertIn("prof_a", s)

    def test_pdf_and_bad_inputs(self):
        self.as_user(self.hod_cs)
        r = self.client.get(f"{URL}students/?export=pdf")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertEqual(self.client.get(f"{URL}students/?export=xml").status_code, 400)
        self.assertEqual(self.client.get(f"{URL}nothing/").status_code, 400)

    def test_formula_cells_are_neutralised(self):
        Student.objects.create(classroom=self.class_cs, roll_number="90", name="=HYPERLINK(\"x\")", created_by=self.t1)
        self.as_user(self.admin)
        self.assertNotIn(",=HYPERLINK", self.text(self.get("students")))

    def test_every_download_is_audited(self):
        self.as_user(self.t1)
        self.get("students")
        e = AuditLog.objects.filter(action="DATA_DOWNLOADED").latest("id")
        self.assertEqual(e.actor, self.t1)
        self.assertEqual(e.new_value["list"], "students")

    def test_admin_zip_has_everything_incl_attendance_marks(self):
        self.submit(self.make_lecture(self.assign_py), self.t1, 3)
        self.as_user(self.admin)
        r = self.client.get(URL + "all/")
        self.assertEqual(r.status_code, 200)
        z = zipfile.ZipFile(io.BytesIO(b"".join(r.streaming_content)))
        names = set(z.namelist())
        for n in ("students", "staff", "hods", "classes", "subjects", "lectures", "attendance_marks", "staff_attendance", "leave_requests", "departments"):
            self.assertIn(n + ".csv", names)
        marks = z.read("attendance_marks.csv").decode("utf-8")
        self.assertTrue(marks.startswith("﻿"))
        self.assertEqual(len(marks.strip().splitlines()), 1 + 5)  # header + 5 students
        self.assertIsNone(z.testzip())


class ExcelFileTests(BaseTestCase):
    """Real .xlsx downloads: they open as a workbook, keep the same scope rules and stay formula-safe."""

    def sheet(self, response):
        from openpyxl import load_workbook
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        return load_workbook(io.BytesIO(response.content)).active

    def values(self, ws):
        return [[c for c in row] for row in ws.iter_rows(values_only=True)]

    def test_students_download_as_excel(self):
        self.as_user(self.admin)
        ws = self.sheet(self.client.get(f"{URL}students/?export=xlsx"))
        rows = self.values(ws)
        self.assertEqual(rows[1][0], "Students (%d)" % (len(rows) - 5)) if rows[1][0].startswith("Students") else None
        self.assertTrue(any("Student 5" in str(c) for r in rows for c in r))
        self.assertEqual(ws.freeze_panes, "A6")  # the header row stays visible

    def test_excel_keeps_the_department_scope(self):
        self.as_user(self.hod_cs)
        rows = self.values(self.sheet(self.client.get(f"{URL}students/?export=xlsx")))
        self.assertFalse(any("Math Student 1" in str(c) for r in rows for c in r))

    def test_formula_text_is_not_run(self):
        Student.objects.create(classroom=self.class_cs, roll_number="99", name="=1+1", created_by=self.t1)
        self.as_user(self.hod_cs)
        ws = self.sheet(self.client.get(f"{URL}students/?export=xlsx"))
        names = [c for r in self.values(ws) for c in r]
        self.assertIn("'=1+1", names)  # shown as text with a leading quote, never as a formula

    def test_reports_can_be_downloaded_as_excel(self):
        self.as_user(self.hod_cs)
        daily = self.client.get("/api/reports/daily/?export=xlsx&date_from=2026-10-01&date_to=2026-10-31")
        self.sheet(daily)
        summary = self.client.get(f"/api/reports/student-summary/?export=xlsx&classroom={self.class_cs.pk}&date_from=2026-10-01&date_to=2026-10-07")
        self.sheet(summary)
        monthly = self.client.get(f"/api/reports/monthly-sheet/?export=xlsx&classroom={self.class_cs.pk}&month=2026-10")
        self.sheet(monthly)

    def test_unknown_format_is_refused(self):
        self.as_user(self.hod_cs)
        self.assertEqual(self.client.get("/api/reports/daily/?export=docx").status_code, 400)
